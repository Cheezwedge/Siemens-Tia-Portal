"""Open the generated workbooks with openpyxl - a reader this repo did not write.

The workbook reader and writer in tiagen are both this repo's code, so a round trip
between them proves they agree with each other, not that either matches the file
format. This checks the writer against an independent implementation, with every
warning it raises treated as a failure. It runs in CI only; openpyxl is never a
dependency of tiagen itself.

    pip install openpyxl
    python tools/check_workbook_independently.py
"""

import os
import subprocess
import sys
import tempfile
import warnings

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main() -> int:
    import openpyxl

    warnings.simplefilter("error")
    work = tempfile.mkdtemp(prefix="tiagen-xcheck-")
    env = dict(os.environ, PYTHONPATH=os.path.join(REPO, "generator"))
    books = {"template": os.path.join(work, "template.xlsx")}
    subprocess.check_call([sys.executable, "-m", "tiagen", "excel", books["template"]], env=env,
                          stdout=subprocess.DEVNULL)
    for example in ("minimal", "conveyor-line", "step-sequence"):
        books[example] = os.path.join(work, f"{example}.xlsx")
        subprocess.check_call([sys.executable, "-m", "tiagen", "excel", books[example], "--from",
                               os.path.join(REPO, "spec", "examples", f"{example}.yaml")],
                              env=env, stdout=subprocess.DEVNULL)

    failures = 0
    for name, path in books.items():
        try:
            wb = openpyxl.load_workbook(path)
            expected = ["Machine", "Equipment", "IO devices", "Modules", "Sequence", "Help"]
            assert wb.sheetnames == expected, f"sheets {wb.sheetnames}"
            equipment = wb["Equipment"]
            assert equipment["A1"].value == "Name", "Equipment header missing"
            assert equipment.freeze_panes == "A2", "header row is not frozen"
            assert equipment.data_validations.dataValidation, "no drop-down lists"
            print(f"ok    {name}: {sum(ws.max_row for ws in wb)} rows over {len(wb.sheetnames)} sheets")
        except Exception as exc:
            failures += 1
            print(f"FAIL  {name}: {type(exc).__name__}: {exc}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
